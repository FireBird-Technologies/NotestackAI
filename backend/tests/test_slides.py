"""Slide decks: what the model writes is clamped to a fixed shape, laid out by code in either theme, checked against
the sources, and drawn the same way as pages, a PDF and an editable PowerPoint. The LLM is faked at app.llm.run."""

import copy
import io
import re

import pytest
from pptx import Presentation

from app.infographics.image import chrome
from app.slides import export
from app.slides.build import finish, view
from app.slides.content import LIMITS, clamp_deck, clamp_slide
from app.slides.layout import GEOMETRY, VARIANTS, assign_variants, compose
from app.slides.render import deck_print_html, slide_html
from tests.test_buildout import _notebook_with_pricing, auth, feed, llm  # noqa: F401  (fixtures)

LONG = "Words that keep going and going " * 40

FULL = {
    "title": {"heading": "Pricing that readers say yes to", "kicker": "From 4 posts", "lead": "What two years taught us."},
    "agenda": {"heading": "What we will cover", "points": [{"text": f"Part {i}"} for i in range(9)]},
    "section": {"heading": "The numbers behind trials", "kicker": "Part two", "lead": "Short trials converted twice as well."},
    "points": {"heading": "Why tiers work", "kicker": "Tiers", "lead": "Readers self-select.",
               "points": [{"term": "Anchoring", "text": "A higher tier makes the middle one feel fair."}] * 4},
    "two_column": {"heading": "Monthly or annual", "left": {"label": "Monthly", "items": ["Low commitment", "Easy to leave"]},
                   "right": {"label": "Annual", "items": ["Two months free", "Lower churn"]}},
    "stat": {"heading": "Short trials win", "stat": {"value": "2x", "label": "more paid signups"}, "lead": "A decision point."},
    "quote": {"heading": "In their words", "quote": {"text": "Thursday would feel empty without it.", "by": "A reader"}},
    "closing": {"heading": "Key takeaways", "takeaways": ["One.", "Two.", "Three."], "closing": "Price the habit."},
}


def _stuff(slide: dict, text: str) -> dict:
    """The same slide with every text field set to `text` (the longest anything could be)."""
    s = copy.deepcopy(slide)
    for k in ("heading", "kicker", "lead", "closing"):
        if k in s:
            s[k] = text
    for p in s.get("points") or []:
        p["text"], p["term"] = text, text[:60]
    for side in ("left", "right"):
        if side in s:
            s[side] = {"label": text, "items": [text] * 6}
    if "stat" in s:
        s["stat"] = {"value": "123456789012345678", "label": text}
    if "quote" in s:
        s["quote"] = {"text": text, "by": text}
    if "takeaways" in s:
        s["takeaways"] = [text] * 5
    return s


@pytest.mark.parametrize("theme", ["dark-space", "light-space"])
@pytest.mark.parametrize("fmt", ["detailed", "presenter"])
def test_every_variant_draws_inside_its_boxes_at_any_length(theme, fmt):
    for layout, variants in VARIANTS.items():
        for v in variants:
            for text in ("Short text here", LONG, "長い日本語のテキスト" * 30, "Supercalifragilistic" * 12):
                raw = {**_stuff(FULL[layout], text), "layout": layout, "variant": v}
                deck = clamp_deck({"title": "Deck", "slides": [raw, raw]}, fmt)
                for i in range(2):
                    c = compose(deck, i, theme, fmt, seed=7)
                    for e in c.elements:
                        assert 0 <= e.x and e.x + e.w <= 1920 + 1 and 0 <= e.y and e.y + e.h <= 1080 + 1, (layout, v, e.role)
                        if e.kind == "text":
                            assert e.text_height() <= e.h + 1, (layout, v, e.role, text[:20])


def test_clamp_never_lets_bad_input_through():
    garbage = [None, 5, "text", [], {"layout": "nonsense", "heading": "<script>alert(1)</script>**Bold** [1]"},
               {"layout": "stat", "heading": "No number", "stat": {"value": "", "label": "x"}},
               {"layout": "quote", "heading": "No quote", "quote": None, "points": [{"text": "a"}, {"text": "b"}]},
               {"layout": "two_column", "heading": "One side", "left": {"items": ["x"]}, "right": None},
               {"layout": "points", "heading": "Nul\x00 byte— dash", "points": "not a list"},
               {"layout": "closing", "heading": "", "takeaways": None, "closing": None},
               {"layout": "points", "heading": "x" * 5000, "points": [{"term": None, "text": {"nested": 1}}] * 30}]
    deck = clamp_deck({"title": None, "slides": garbage}, "detailed")
    assert deck and deck["title"]
    keys = set(clamp_slide({"heading": "x"}, "detailed"))
    for s in deck["slides"]:
        assert set(s) == keys  # every slide has every key
        for value in (s["heading"], s["lead"], s["kicker"]):
            assert "<" not in value and "\x00" not in value and "—" not in value and "**" not in value
        assert len(s["heading"]) <= LIMITS["detailed"]["heading"]
    layouts = [s["layout"] for s in deck["slides"]]
    assert layouts[0] == "points" or layouts[0] == "section"  # unknown layout becomes a simpler one
    assert "stat" not in layouts and "quote" not in layouts and "two_column" not in layouts
    assert clamp_deck({"slides": []}, "detailed") is None and clamp_deck("nope", "presenter") is None


def test_pages_escape_everything_the_model_wrote():
    raw = {"layout": "points", "heading": "Fine <b onmouseover=x>heading</b>", "lead": "a & b < c",
           "points": [{"term": "<img src=x>", "text": "\"quoted\" <script>x</script> text"}] * 3}
    deck = clamp_deck({"title": "<title>", "slides": [raw]}, "detailed")
    deck["slides"][0]["variant"] = "a"
    html = slide_html(deck, 0, "dark-space", "detailed", 1)
    assert "<script" not in html and "<img" not in html and "onmouseover" not in html
    assert "a &amp; b &lt; c" in html


def test_variants_never_repeat_in_a_row_and_differ_between_decks():
    slides = [clamp_slide({**FULL[k], "layout": k}, "detailed") for k in
              ("title", "agenda", "points", "points", "points", "stat", "points", "quote", "points", "closing")]
    a, b = copy.deepcopy(slides), copy.deepcopy(slides)
    assign_variants(a, "dark-space", seed=1)
    pairs = [(s["layout"], s["variant"]) for s in a]
    assert all(pairs[i] != pairs[i + 1] for i in range(len(pairs) - 1))
    points = [s["variant"] for s in a if s["layout"] == "points"]
    assert len(set(points)) == len(VARIANTS["points"])  # every points variant used before any repeats
    assign_variants(b, "dark-space", seed=2, previous=pairs)
    assert [s["variant"] for s in b] != [s["variant"] for s in a]
    assert b[0]["variant"] != a[0]["variant"] and b[-1]["variant"] != a[-1]["variant"]  # new opener and closer
    again = copy.deepcopy(slides)
    assign_variants(again, "dark-space", seed=1)
    assert [s["variant"] for s in again] == [s["variant"] for s in a]  # one deck always draws the same


def test_the_two_themes_compose_differently():
    for layout, variants in VARIANTS.items():
        for v in variants:
            assert GEOMETRY["dark-space"][layout][v] != GEOMETRY["light-space"][layout][v] or layout in ("points", "two_column"), \
                (layout, v)


def _deck(fmt: str = "detailed") -> dict:
    src = [{"path": "p.md", "line_start": 1, "line_end": 2, "title": "On Pricing"}]
    slides = [{**FULL[k], "layout": k, "notes": f"Say {k}", "sources": src}
              for k in ("title", "agenda", "points", "two_column", "stat", "quote", "closing")]
    deck = clamp_deck({"title": "Pricing", "slides": slides}, fmt)
    assign_variants(deck["slides"], "light-space", seed=3)
    return deck


def test_powerpoint_has_every_slide_as_editable_text_with_notes():
    deck = _deck()
    data = export.build_pptx(deck, "light-space", "detailed", 3)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(deck["slides"]) and prs.slide_width == 12192000
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Why tiers work" in texts[2] and "Anchoring" in texts[2] and "middle one feel fair" in texts[2]
    notes = prs.slides[2].notes_slide.notes_text_frame.text
    assert "Say points" in notes and "On Pricing (lines 1-2)" in notes


@pytest.mark.skipif(not chrome(), reason="no Chrome to print with")
def test_pdf_has_one_page_per_slide_and_nothing_overflows():
    from pypdf import PdfReader

    deck = finish(_deck(), "dark-space", "detailed", 3)
    assert export.overflowing(deck, "dark-space", "detailed", 3) == set()
    pdf = PdfReader(io.BytesIO(export.render_pdf(deck, "dark-space", "detailed", 3)))
    assert len(pdf.pages) == len(deck["slides"])
    assert round(float(pdf.pages[0].mediabox.width) / float(pdf.pages[0].mediabox.height), 2) == round(1920 / 1080, 2)


def test_a_stored_deck_with_a_broken_slide_still_shows():
    content = {"theme": "dark-space", "format": "detailed", "seed": "abc", "deck": _deck()}
    content["deck"]["slides"][2]["variant"] = "zz"  # unknown: drawn with the first variant
    content["deck"]["slides"][3]["points"] = None
    pages = view(content)
    assert len(pages) == len(content["deck"]["slides"]) and all(p.startswith("<!doctype html>") for p in pages)
    assert view({"deck": {"slides": []}}) == [] and view({}) == []
    assert "Pricing" in deck_print_html(clamp_deck(content["deck"], "detailed"), "x", "detailed", 1)


# --- the whole flow, through the API ------------------------------------------------------------------------------

@pytest.fixture()
def no_chrome(monkeypatch):
    """Generation measures slides in Chrome when one is installed; the API tests run without it."""
    monkeypatch.setattr(export, "chrome", lambda: None)


def _fake_many(monkeypatch, pricing_path: str, log: dict):
    from app.llm import run

    def predict_many(signature, inputs_list, *, db, workspace_id, lm=None, workers=5):
        name = signature.__name__
        log.setdefault(name, []).extend(inputs_list)
        out = []
        for inp in inputs_list:
            if name == "WriteSlide":
                good = {"path": pricing_path, "line_start": 1, "line_end": 12}
                fake = {"path": "sources/made-up.md", "line_start": 1, "line_end": 3}
                slide = {"kicker": "Pricing", "heading": inp["heading"], "lead": "Charging more made readers take it seriously.",
                         "points": [{"term": "Price", "text": "The paid tier doubled after the price went up."},
                                    {"term": "Invented", "text": "Ninety percent of readers loved it."},
                                    {"term": "Seriousness", "text": "Readers took the work seriously."}],
                         "notes": "Talk about pricing.", "sources": [good, fake]}
                if inp["layout"] == "stat":
                    slide["stat"] = {"value": "90%", "label": "of readers upgraded"}  # not in the post: removed
                if inp["layout"] == "quote":
                    slide["quote"] = {"text": "Charging more made my readers take the work seriously.", "by": "Ada"}
                if inp["layout"] == "two_column":
                    slide["left"] = {"label": "Before", "items": ["Low price", "Few paid readers"]}
                    slide["right"] = {"label": "After", "items": ["Ten dollars", "Paid tier doubled"]}
                out.append({"slide": slide})
            elif name == "CheckSlide":
                bad = [c for c in inp["claims"] if "Ninety percent" in c]
                out.append({"checks": [{"n": int(c.split(".")[0]), "supported": False, "fix": ""} for c in bad]})
            else:
                out.append(None)
        return out

    monkeypatch.setattr(run, "predict_many", predict_many)


def _outline(n: int, layouts=("points", "quote", "stat", "points", "section", "points", "points", "two_column", "points")):
    def make(**inputs):
        return {"outline": {"title": "Pricing for writers", "subtitle": "Why charging more worked",
                            "opening_kicker": "Pricing", "agenda_heading": "What we cover",
                            "slides": [{"layout": layouts[i % len(layouts)], "heading": f"Slide idea {i}",
                                        "purpose": "Make a point", "sources": []} for i in range(n)]}}
    return make


CLOSING = {"closing": {"kicker": "Conclusion", "heading": "Key takeaways",
                       "takeaways": ["Charge more.", "Readers take it seriously.", "The paid tier grows."],
                       "closing": "Price the work like it matters.", "notes": "Close."}}


def test_short_detailed_deck_is_grounded_checked_and_downloadable(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    log: dict = {}
    _fake_many(monkeypatch, pricing["path"], log)
    llm["OutlineDeck"] = _outline(9)  # too many for a short deck: trimmed
    llm["WriteClosing"] = CLOSING
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_format": "detailed", "deck_length": "short",
                      "theme": "light-space", "instructions": "For new writers"}, headers=auth).json()
    assert art["type_label"] == "Slide deck"
    run_jobs()
    done = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()
    assert done["status"] == "ready", done["content"].get("error")
    c = done["content"]
    slides = c["deck"]["slides"]
    assert 6 <= len(slides) <= 8 and len(done["content"]["slides_html"]) == len(slides)
    assert slides[0]["layout"] == "title" and slides[-1]["layout"] == "closing"
    assert slides[-1]["closing"] == "Price the work like it matters." and len(slides[-1]["takeaways"]) == 3
    assert not any(s["layout"] == "agenda" for s in slides)  # short decks have no agenda
    assert c["theme"] == "light-space" and c["format"] == "detailed"
    body = [s for s in slides[1:-1]]
    for s in body:
        assert all(src["path"] == pricing["path"] for src in s["sources"])  # the invented path is gone
        assert all("Ninety percent" not in p["text"] for p in s["points"])  # unsupported claim removed
        assert s["layout"] != "stat"  # its number was not in the post
    assert any(s["layout"] == "quote" for s in body)  # the quote is in the post word for word
    pairs = [(s["layout"], s["variant"]) for s in slides]
    assert all(pairs[i] != pairs[i + 1] for i in range(len(pairs) - 1))
    assert all(log["WriteSlide"][0]["source_lines"]) and log["WriteSlide"][0]["deck_format"] == "detailed"
    html = "".join(c["slides_html"])
    assert "<script" not in html and "Pricing for writers" in html

    pptx = client.get(f"/api/artifacts/{art['id']}/slides.pptx", headers=auth)
    assert pptx.status_code == 200 and "Pricing for writers.pptx" in pptx.headers["content-disposition"]
    assert len(Presentation(io.BytesIO(pptx.content)).slides) == len(slides)
    assert client.get(f"/api/artifacts/{art['id']}/slides.pdf", headers=auth).status_code == 503  # no Chrome here
    assert client.get(f"/api/artifacts/{art['id']}/slides.doc", headers=auth).status_code == 422


def test_default_presenter_deck_has_an_agenda_and_stays_in_range(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    log: dict = {}
    _fake_many(monkeypatch, pricing["path"], log)
    llm["OutlineDeck"] = _outline(7)
    llm["WriteClosing"] = CLOSING
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_format": "presenter", "deck_length": "default",
                      "language": "Spanish"}, headers=auth).json()
    run_jobs()
    done = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()
    slides = done["content"]["deck"]["slides"]
    assert done["status"] == "ready" and 9 <= len(slides) <= 11
    assert slides[1]["layout"] == "agenda" and [p["text"] for p in slides[1]["points"]][:2] == ["Slide idea 0", "Slide idea 1"]
    assert done["content"]["theme"] == "dark-space"  # the default
    assert log["WriteSlide"][0]["language"] == "Spanish" and log["WriteSlide"][0]["deck_format"] == "presenter"
    for s in slides:
        assert all(len(p["text"]) <= LIMITS["presenter"]["text"] for p in s["points"] if s["layout"] != "agenda")


def test_too_little_material_fails_cleanly_and_sources_are_required(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    _fake_many(monkeypatch, pricing["path"], {})
    llm["OutlineDeck"] = _outline(1)
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "default"}, headers=auth).json()
    run_jobs()
    failed = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()
    assert failed["status"] == "failed" and "Short" in failed["content"]["error"]
    assert llm["_calls"].count("OutlineDeck") >= 2  # asked again before giving up
    assert client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                       "document_ids": []}, headers=auth).status_code == 400
    assert client.get(f"/api/artifacts/{art['id']}/slides.pptx", headers=auth).status_code == 404


def test_a_slide_the_model_fails_to_write_is_written_from_its_sources(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    from app.llm import run

    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    log: dict = {}
    _fake_many(monkeypatch, pricing["path"], log)
    real = run.predict_many

    def flaky(signature, inputs_list, **kw):
        out = real(signature, inputs_list, **kw)
        return [None] * len(out) if signature.__name__ == "WriteSlide" else out

    monkeypatch.setattr(run, "predict_many", flaky)
    llm["OutlineDeck"] = lambda **kw: {"outline": {"title": "Pricing", "subtitle": "", "opening_kicker": "", "agenda_heading": "",
                                                   "slides": [{"layout": "points", "heading": f"Idea {i}", "purpose": "A point",
                                                               "sources": [{"path": pricing["path"], "line_start": 1, "line_end": 14}]}
                                                              for i in range(4)]}}
    llm["WriteClosing"] = lambda **kw: (_ for _ in ()).throw(RuntimeError("model down"))
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "short"}, headers=auth).json()
    run_jobs()
    done = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()
    assert done["status"] == "ready", done["content"].get("error")
    slides = done["content"]["deck"]["slides"]
    assert [s["heading"] for s in slides[1:-1]] == [f"Idea {i}" for i in range(4)]
    assert slides[-1]["layout"] == "closing" and slides[-1]["takeaways"]
    assert len(log["WriteSlide"]) == 8  # each slide tried twice
    assert re.search(r"Charging more|ten dollars|seriously", " ".join(str(s) for s in slides[1:-1]))


def test_the_opening_and_closing_say_only_what_the_slides_say(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    _fake_many(monkeypatch, pricing["path"], {})
    outline = _outline(4)

    def mixed_up(**kw):
        out = outline(**kw)
        out["outline"]["subtitle"] = "Ada won a writing award for her pricing"  # joins two things no slide says
        return out

    llm["OutlineDeck"] = mixed_up
    takeaways = ["Charge more.", "Ada sold a million copies.", "The paid tier grows."]
    llm["WriteClosing"] = {"closing": {**CLOSING["closing"], "takeaways": takeaways}}

    def check(**kw):
        claims = kw["claims"]
        return {"checks": [{"n": int(c.split(".")[0]), "supported": False, "fix": ""} for c in claims
                           if "award" in c or "million" in c]}

    llm["CheckSlide"] = check
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "short"}, headers=auth).json()
    run_jobs()
    slides = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()["content"]["deck"]["slides"]
    assert "award" not in slides[0]["lead"] and slides[0]["heading"] == "Pricing for writers"
    assert slides[-1]["takeaways"] == ["Charge more.", "The paid tier grows."]


def test_long_deck_has_parts_and_the_request_reaches_every_step(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    log: dict = {}
    _fake_many(monkeypatch, pricing["path"], log)
    layouts = ("section", "points", "stat", "points", "section", "quote", "points", "two_column", "section", "points",
               "points", "points")
    llm["OutlineDeck"] = _outline(12, layouts)
    llm["CheckRequest"] = {"asks": [{"ask": "For new writers", "covered": True}]}
    seen: dict = {}

    def closing(**kw):
        seen["closing"] = kw["request"]
        return CLOSING

    llm["WriteClosing"] = closing
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "long", "instructions": "For new writers"},
                      headers=auth).json()
    run_jobs()
    done = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()
    assert done["status"] == "ready", done["content"].get("error")
    slides = done["content"]["deck"]["slides"]
    assert 12 <= len(slides) <= 16 and done["content"]["length"] == "long"
    sections = [s["heading"] for s in slides if s["layout"] == "section"]
    assert slides[1]["layout"] == "agenda" and [p["text"] for p in slides[1]["points"]] == sections  # the parts
    assert all(w["request"] == "For new writers" for w in log["WriteSlide"]) and seen["closing"] == "For new writers"
    assert llm["_calls"].count("OutlineDeck") == 1  # nothing missed: no second plan


def test_a_missed_ask_is_planned_again(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    _fake_many(monkeypatch, pricing["path"], {})
    requests: list[str] = []
    first, second = _outline(4), _outline(5)

    def outline(**kw):
        requests.append(kw["request"])
        out = (first if len(requests) == 1 else second)(**kw)
        if len(requests) > 1:
            out["outline"]["slides"][-1]["heading"] = "Pricing for beginners"
        return out

    llm["OutlineDeck"] = outline
    llm["CheckRequest"] = {"asks": [{"ask": "a slide for beginners", "covered": False}, {"ask": "pricing", "covered": True}]}
    llm["WriteClosing"] = CLOSING
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "short",
                      "instructions": "Pricing, with a slide for beginners"}, headers=auth).json()
    run_jobs()
    slides = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()["content"]["deck"]["slides"]
    assert len(requests) == 2 and "a slide for beginners" in requests[1] and "pricing" not in requests[1].split("supports: ")[1]
    assert any(s["heading"] == "Pricing for beginners" for s in slides)
    assert client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                       "document_ids": [pricing["id"]], "deck_length": "huge"}, headers=auth).status_code == 422


def test_the_closing_is_tried_again_and_a_repeated_heading_goes(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    _fake_many(monkeypatch, pricing["path"], {})
    outline = _outline(6, ("points", "section", "points", "quote", "points", "two_column"))

    def with_repeat(**kw):
        out = outline(**kw)
        out["outline"]["slides"][1]["heading"] = out["outline"]["slides"][2]["heading"]  # a part titled like its slide
        return out

    llm["OutlineDeck"] = with_repeat
    tries = []

    def closing(**kw):
        tries.append(1)
        if len(tries) == 1:
            raise RuntimeError("model hiccup")
        return CLOSING

    llm["WriteClosing"] = closing
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "short"}, headers=auth).json()
    run_jobs()
    slides = client.get(f"/api/artifacts/{art['id']}", headers=auth).json()["content"]["deck"]["slides"]
    assert len(tries) == 2 and slides[-1]["closing"] == "Price the work like it matters."
    headings = [s["heading"] for s in slides]
    assert len(headings) == len(set(headings)) and 6 <= len(slides) <= 8


@pytest.mark.parametrize("key", [("title", "a"), ("title", "c"), ("section", "a"), ("section", "b")])
def test_light_planets_sit_on_their_ring(key):
    """A planet beside rings is drawn on one of them, tilt included, whatever the seed."""
    import math

    from app.slides.decor import build

    for seed in range(40):
        shapes = build("light-space", *key, seed, 0).shapes
        rings = [s for s in shapes if s.kind == "ring"]
        for p in (s for s in shapes if s.kind == "planet"):
            def off(r, p=p):  # how far the planet's centre is from the ring's line, in units of the ring
                rot = math.radians(-r.rot)
                dx, dy = p.cx - r.cx, p.cy - r.cy
                x, y = dx * math.cos(rot) - dy * math.sin(rot), dx * math.sin(rot) + dy * math.cos(rot)
                return abs(math.hypot(x / r.rx, y / r.ry) - 1)
            assert min(off(r) for r in rings) < 1e-6, (key, seed)


@pytest.mark.parametrize("key", [("title", "b"), ("points", "a")])
def test_dark_rings_wrap_round_their_planet(key):
    """A ringed planet's far half is drawn before the planet (hidden behind it) and the near half after (in front)."""
    from app.slides.decor import FAR, NEAR, build

    for seed in range(10):
        shapes = build("dark-space", *key, seed, 0).shapes
        at = next(i for i, s in enumerate(shapes) if s.kind == "planet")
        rings = [(i, s) for i, s in enumerate(shapes) if s.kind in ("ring", "band")]
        assert rings and all(s.arc for _, s in rings)
        assert all(i < at for i, s in rings if s.arc == FAR) and all(i > at for i, s in rings if s.arc == NEAR)
        assert all(s.rx * s.inner > shapes[at].r for _, s in rings if s.kind == "band")  # the band clears the planet


@pytest.mark.parametrize("theme", ["dark-space", "light-space"])
def test_every_variant_marks_its_editable_text(theme):
    """Every piece of a slide's own text is drawn in a span the editor can find, with the field's real limit; numbers
    and footers are not editable."""
    from app.slides.content import field_limit

    for fmt in ("detailed", "presenter"):
        for layout, variants in GEOMETRY[theme].items():
            for v in variants:
                deck = clamp_deck({"title": "T", "slides": [{"layout": layout, **FULL[layout]}]}, fmt)
                deck["slides"][0]["variant"] = v
                html = slide_html(deck, 0, theme, fmt, 3)
                found = dict(re.findall(r'data-f="([^"]+)" data-max="(\d+)"', html))
                assert "heading" in found or layout in ("quote", "closing"), (layout, v)
                assert all(int(n) == field_limit(layout, f, fmt) for f, n in found.items())
                s = deck["slides"][0]
                want = {"points": [f"points.{i}.text" for i in range(len(s["points"]))], "stat": ["stat.value", "stat.label"],
                        "quote": ["quote.text", "quote.by"], "two_column": ["left.label", "right.items.1"],
                        "closing": [f"takeaways.{i}" for i in range(len(s["takeaways"]))] + ["closing"]}.get(layout, [])
                assert set(want) <= set(found), (fmt, layout, v, found)
                assert not any(f.startswith("footer") for f in found)


def _ready_deck(client, auth, run_jobs, monkeypatch, llm):  # noqa: F811
    pricing, nb = _notebook_with_pricing(client, auth, run_jobs)
    _fake_many(monkeypatch, pricing["path"], {})
    llm["OutlineDeck"] = _outline(5)
    llm["WriteClosing"] = CLOSING
    art = client.post("/api/artifacts/generate", json={"type": "slide_deck", "notebook_id": nb["id"],
                      "document_ids": [pricing["id"]], "deck_length": "short"}, headers=auth).json()
    run_jobs()
    return client.get(f"/api/artifacts/{art['id']}", headers=auth).json()


def test_edited_deck_is_saved_fitted_and_downloaded(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    done = _ready_deck(client, auth, run_jobs, monkeypatch, llm)
    slides = copy.deepcopy(done["content"]["deck"]["slides"])
    assert len(done["content"]["slide_slots"]) == len(slides) and "heading" in done["content"]["slide_slots"][0]
    pi = next(i for i, s in enumerate(slides) if s["points"] and s["layout"] == "points")
    slides[0]["heading"] = "A brand new title"
    slides[0]["notes"] = "Say hello first."
    slides[pi]["points"][0]["text"] = "Edited point " + "very long " * 60
    slides[-1]["takeaways"][1] = "  An <b>edited</b> takeaway  "
    url = f"/api/artifacts/{done['id']}/deck"
    saved = client.put(url, json={"slides": slides}, headers=auth)
    assert saved.status_code == 200, saved.text
    after = saved.json()["content"]
    s = after["deck"]["slides"]
    assert s[0]["heading"] == "A brand new title" and s[0]["notes"] == "Say hello first."
    assert s[pi]["points"][0]["text"].startswith("Edited point") and len(s[pi]["points"][0]["text"]) <= LIMITS["detailed"]["text"]
    assert s[-1]["takeaways"][1] == "An edited takeaway"  # plain text only
    before = done["content"]["deck"]["slides"]
    assert [(x["layout"], x["variant"]) for x in s] == [(x["layout"], x["variant"]) for x in before]  # designs kept
    assert "A brand new title" in after["slides_html"][0]
    pptx = client.get(f"/api/artifacts/{done['id']}/slides.pptx", headers=auth)
    texts = [sh.text_frame.text for sh in Presentation(io.BytesIO(pptx.content)).slides[0].shapes if sh.has_text_frame]
    assert any("A brand new title" in t for t in texts)

    assert client.put(url, json={"slides": []}, headers=auth).status_code == 422
    assert client.put(url, json={"slides": [{"layout": "points", "heading": ""}]}, headers=auth).status_code == 400
    assert client.put(url, json={"slides": [slides[0]] * 31}, headers=auth).status_code == 400
    kept = client.get(f"/api/artifacts/{done['id']}", headers=auth).json()["content"]["deck"]["slides"]
    assert kept[0]["heading"] == "A brand new title" and len(kept) == len(slides)  # a refused save changes nothing


def test_slides_can_be_added_removed_moved_and_change_layout(client, auth, run_jobs, feed, llm, monkeypatch, no_chrome):  # noqa: F811
    done = _ready_deck(client, auth, run_jobs, monkeypatch, llm)
    slides = copy.deepcopy(done["content"]["deck"]["slides"])
    n = len(slides)
    pi = next(i for i, s in enumerate(slides) if s["layout"] == "points")
    edited = copy.deepcopy(slides)
    edited[pi]["points"].append({"term": "", "text": "A point added by hand"})  # a text box added
    edited[pi]["lead"] = ""  # and one removed
    stat = {"layout": "stat", "variant": "", "heading": "A new stat slide", "stat": {"value": "42%", "label": "Readers who stayed"}}
    edited.insert(1, stat)  # a slide added
    edited.insert(2, copy.deepcopy(edited[0]) | {"variant": ""})  # a slide duplicated
    del edited[-2]  # a slide removed
    edited[-1], edited[-2] = edited[-2], edited[-1]  # two moved
    url = f"/api/artifacts/{done['id']}"
    prev = client.post(f"{url}/deck-preview", json={"slides": edited}, headers=auth)
    assert prev.status_code == 200, prev.text
    got = prev.json()
    out = got["slides"]
    assert len(out) == n + 1 and len(got["slides_html"]) == n + 1 and len(got["slide_slots"]) == n + 1
    assert out[1]["layout"] == "stat" and out[1]["variant"] and "42%" in got["slides_html"][1]
    assert "value" in got["slide_slots"][1] and "label" in got["slide_slots"][1]
    moved = next(s for s in out if s["heading"] == slides[pi]["heading"])
    assert moved["points"][-1]["text"] == "A point added by hand" and moved["lead"] == ""
    pairs = [(x["layout"], x["variant"]) for x in out]
    assert all(a != b for a, b in zip(pairs, pairs[1:], strict=False))  # no two alike in a row
    assert client.get(url, headers=auth).json()["content"]["deck"]["slides"] == done["content"]["deck"]["slides"]  # preview saves nothing

    # a layout change: the points slide becomes a quote slide
    q = next(i for i, s in enumerate(edited) if s["layout"] == "points")
    edited[q] = {**edited[q], "layout": "quote", "variant": "", "quote": {"text": "Price the work like it matters.", "by": "A writer"}}
    saved = client.put(f"{url}/deck", json={"slides": edited}, headers=auth).json()["content"]
    assert saved["deck"]["slides"][q]["layout"] == "quote" and saved["deck"]["slides"][q]["variant"]
    assert len(saved["deck"]["slides"]) == n + 1
    pptx = client.get(f"{url}/slides.pptx", headers=auth)
    assert len(Presentation(io.BytesIO(pptx.content)).slides) == n + 1


def test_fill_variants_keeps_good_variants_and_fixes_the_rest():
    from app.slides.build import fill_variants

    deck = clamp_deck({"title": "T", "slides": [{"layout": lay, **FULL[lay]} for lay in ("title", "points", "points", "stat", "closing")]},
                      "detailed")
    assign_variants(deck["slides"], "dark-space", seed=3)
    before = [s["variant"] for s in deck["slides"]]
    fill_variants(deck["slides"], "dark-space", 3)
    assert [s["variant"] for s in deck["slides"]] == before  # nothing to fix: nothing changes
    deck["slides"][3]["variant"] = ""  # new
    deck["slides"][2]["variant"] = deck["slides"][1]["variant"]  # the same as the slide before it
    fill_variants(deck["slides"], "dark-space", 3)
    assert deck["slides"][3]["variant"] in VARIANTS["stat"]
    assert deck["slides"][2]["variant"] != deck["slides"][1]["variant"]
    assert [s["variant"] for s in deck["slides"]][:2] == before[:2] and deck["slides"][4]["variant"] == before[4]
